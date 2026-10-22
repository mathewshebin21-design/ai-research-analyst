import io

from pptx import Presentation

from tests.conftest import login
from tests.test_api import _completed_run


def test_share_link_is_read_only_public_and_revocable(client):
    ana = login(client)
    ws, d = _completed_run(client, ana)
    r = client.post(f"/runs/{d['id']}/shares", headers=ana)
    assert r.status_code == 201
    token, share_id = r.json()["token"], r.json()["id"]
    assert "token" not in client.get(f"/runs/{d['id']}/shares", headers=ana).json()[0]  # never shown again

    public = client.get(f"/shared/{token}")  # no sign-in
    assert public.status_code == 200
    body = public.json()
    assert body["shared"] is True and body["question"] == d["question"] and body["claims"]
    for private in ("events", "document_ids", "workspace_id"):
        assert private not in body
    assert set(body["usage"]) == {"searches"}  # no costs or token counts
    assert "dev-" not in public.text  # no user identifiers leak

    assert client.get("/shared/not-a-real-token").status_code == 404
    assert client.delete(f"/runs/{d['id']}/shares/{share_id}", headers=ana).json() == {"revoked": True}
    assert client.get(f"/shared/{token}").status_code == 404


def test_only_editors_create_links_and_only_completed_runs_can_be_shared(client):
    ana, ben = login(client, "ana@example.com"), login(client, "ben@example.com")
    ws, d = _completed_run(client, ana)
    client.post(f"/workspaces/{ws['id']}/members", json={"email": "ben@example.com", "role": "viewer"}, headers=ana)
    assert client.post(f"/runs/{d['id']}/shares", headers=ben).status_code == 403
    planned = client.post(
        f"/workspaces/{ws['id']}/runs", json={"question": "Should we enter the Irish market next?"}, headers=ana
    ).json()
    assert client.post(f"/runs/{planned['id']}/shares", headers=ana).status_code == 409


def test_slide_deck_export(client):
    h = login(client)
    ws, d = _completed_run(client, h, currency="USD")
    r = client.post(f"/runs/{d['id']}/slides", headers=h)
    assert r.status_code == 200 and r.headers["content-disposition"].endswith('.pptx"')
    prs = Presentation(io.BytesIO(r.content))
    texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame]
    assert len(prs.slides) >= 8
    assert any(d["recommendation"]["headline"] in t for t in texts)
    assert any("Not financial advice" in t for t in texts)
    charts = [sh for s in prs.slides for sh in s.shapes if sh.has_chart]
    assert charts and len(charts[0].chart.series) == 3
    tables = [sh.table for s in prs.slides for sh in s.shapes if sh.has_table]
    assert any("$" in c.text for t in tables for row in t.rows for c in row.cells)
