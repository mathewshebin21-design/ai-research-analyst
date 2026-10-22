"""Slide-deck export (16:9) for presenting a run. Same content rules as the PDF: every figure is a
claim, an assumption or a model output, and the deck ends with the not-advice notice."""

import io
from datetime import UTC, datetime

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from intel.report_pdf import DISCLAIMER, STATUS_LABEL

NAVY = RGBColor(0x0B, 0x16, 0x28)
IVORY = RGBColor(0xEF, 0xE9, 0xDC)
GOLD = RGBColor(0xC8, 0xA6, 0x5D)
INK = RGBColor(0x1B, 0x23, 0x33)
MUTED = RGBColor(0x5E, 0x66, 0x75)
SERIES = {"Base": GOLD, "Optimistic": RGBColor(0x2C, 0x6E, 0x4F), "Conservative": RGBColor(0x7D, 0x8A, 0xA0)}
W, H = Inches(13.333), Inches(7.5)


def _money(v: float, sym: str) -> str:
    sign = "−" if v < 0 else ""
    v = abs(v)
    return f"{sign}{sym}{v / 1_000_000:.2f}m" if v >= 1_000_000 else f"{sign}{sym}{round(v / 1000):,}k"


class Deck:
    def __init__(self, footer: str) -> None:
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.footer = footer

    def slide(self, title: str, eyebrow: str = "", dark: bool = False):
        s = self.prs.slides.add_slide(self.blank)
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = NAVY if dark else RGBColor(0xFF, 0xFF, 0xFF)
        bar = s.shapes.add_shape(1, Inches(0.6), Inches(0.55), Inches(0.08), Inches(0.9))
        bar.fill.solid()
        bar.fill.fore_color.rgb = GOLD
        bar.line.fill.background()
        if eyebrow:
            self.text(s, eyebrow.upper(), Inches(0.85), Inches(0.5), Inches(11), Inches(0.35), 11, GOLD, bold=True)
        self.text(
            s, title, Inches(0.85), Inches(0.8), Inches(11.8), Inches(0.9), 28, IVORY if dark else NAVY, font="Georgia"
        )
        self.text(s, self.footer, Inches(0.6), Inches(7.0), Inches(12), Inches(0.3), 9, MUTED)
        return s

    @staticmethod
    def text(s, text, x, y, w, h, size, color, bold=False, font="Calibri", bullets: list[str] | None = None):
        box = s.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        lines = bullets if bullets is not None else [text]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = ("• " if bullets is not None else "") + line
            p.space_after = Pt(6)
            for r in p.runs:
                r.font.size, r.font.bold, r.font.name = Pt(size), bold, font
                r.font.color.rgb = color
        return box

    @staticmethod
    def table(s, rows: list[list[str]], x, y, w, col_widths: list[float], size: int = 12):
        shape = s.shapes.add_table(len(rows), len(rows[0]), x, y, w, Inches(0.4) * len(rows))
        t = shape.table
        for j, cw in enumerate(col_widths):
            t.columns[j].width = Emu(int(w * cw))
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                cell.text = str(val)
                cell.fill.solid()
                cell.fill.fore_color.rgb = (
                    NAVY if i == 0 else (RGBColor(0xF7, 0xF5, 0xEF) if i % 2 else RGBColor(0xFF, 0xFF, 0xFF))
                )
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(size)
                        r.font.bold = i == 0
                        r.font.color.rgb = IVORY if i == 0 else INK
        return t


def build_deck(data: dict) -> bytes:
    sym = data.get("currency_symbol", "€")
    rec = data.get("recommendation") or {}
    decision = data.get("decision") or {}
    deck = Deck(f"{data['question'][:110]}  ·  {datetime.now(UTC):%d %b %Y}  ·  Not financial advice")

    # 1. Title
    s = deck.slide(data["question"], "Market entry assessment", dark=True)
    deck.text(
        s,
        f"{data.get('workspace_name', '')} · completed {(data.get('finished_at') or '')[:10]}",
        Inches(0.85),
        Inches(2.4),
        Inches(11),
        Inches(0.5),
        16,
        GOLD,
    )
    counts: dict[str, int] = {}
    for c in data.get("claims", []):
        counts[c["status"]] = counts.get(c["status"], 0) + 1
    deck.text(
        s,
        f"{len(data.get('sources', []))} sources · {len(data.get('claims', []))} claims checked · "
        f"{counts.get('supported', 0)} supported",
        Inches(0.85),
        Inches(3.0),
        Inches(11),
        Inches(0.5),
        14,
        IVORY,
    )

    # 2. Recommendation
    if rec:
        s = deck.slide(rec.get("headline", ""), f"Recommendation · {rec.get('verdict', '')}")
        deck.text(s, rec.get("summary", ""), Inches(0.85), Inches(1.9), Inches(7.4), Inches(3), 16, INK)
        lines = [
            f"Score {rec.get('overall')} / 100",
            f"Confidence {rec.get('confidence_label')} ({rec.get('confidence', 0):.2f})",
        ]
        if decision.get("changed"):
            lines.append(f"Current verdict: {decision['verdict']} (score {decision['overall']})")
        deck.text(s, "", Inches(8.8), Inches(1.9), Inches(4), Inches(2.5), 16, NAVY, bullets=lines)
        deck.text(s, rec.get("rule", ""), Inches(0.85), Inches(5.6), Inches(11.5), Inches(0.8), 11, MUTED)

        s = deck.slide("What has to be true", "Conditions")
        tracked = data.get("conditions") or [{"text": t, "status": "open"} for t in rec.get("conditions", [])]
        labels = {"open": "Open", "met": "Met", "not_met": "Not met"}
        deck.text(
            s,
            "",
            Inches(0.85),
            Inches(1.9),
            Inches(6.2),
            Inches(4.5),
            15,
            INK,
            bullets=[f"{c['text']}  [{labels.get(c['status'], 'Open')}]" for c in tracked],
        )
        deck.text(s, "WHAT WOULD CHANGE THIS", Inches(7.4), Inches(1.9), Inches(5), Inches(0.4), 11, GOLD, bold=True)
        deck.text(s, "", Inches(7.4), Inches(2.3), Inches(5.3), Inches(4), 15, INK, bullets=rec.get("would_change", []))

    # 3. Financial scenarios with a native chart
    scenario = data.get("scenario")
    if scenario:
        s = deck.slide("Cumulative cash over 36 months", "Financial scenarios")
        res = scenario["results"]["scenarios"]
        cd = CategoryChartData()
        cd.categories = [f"M{i + 1}" for i in range(len(res["base"]["cumulative_cash"]))]
        for name, key in (("Optimistic", "optimistic"), ("Base", "base"), ("Conservative", "conservative")):
            cd.add_series(name, [round(v) for v in res[key]["cumulative_cash"]])
        chart = s.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(0.6), Inches(1.8), Inches(8), Inches(4.9), cd).chart
        chart.has_legend, chart.legend.position, chart.legend.include_in_layout = True, XL_LEGEND_POSITION.BOTTOM, False
        for series in chart.series:
            series.format.line.color.rgb = SERIES[series.name]
            series.smooth = False
        axis_sym = "INR " if sym == "₹" else sym
        values = chart.value_axis.tick_labels
        values.number_format, values.number_format_is_linked = f'"{axis_sym}"#,##0,"k";-"{axis_sym}"#,##0,"k"', False
        values.font.size = Pt(10)
        chart.category_axis.tick_labels.font.size = Pt(10)
        chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        cat_ax = chart.category_axis._element
        skip = cat_ax.makeelement(qn("c:tickLblSkip"), {"val": "6"})
        after = cat_ax.find(qn("c:noMultiLvlLbl"))  # schema order: tickLblSkip precedes noMultiLvlLbl
        after.addprevious(skip) if after is not None else cat_ax.append(skip)
        base = res["base"]
        deck.table(
            s,
            [
                ["Base case", ""],
                [
                    "Break-even month",
                    f"Month {base['break_even_month']}" if base["break_even_month"] else "Not in 36 months",
                ],
                ["Peak funding need", _money(-base["peak_funding"], sym)],
                ["Cash at month 36", _money(base["cash_36"], sym)],
                ["Revenue, 36 months", _money(base["revenue_36"], sym)],
            ],
            Inches(8.9),
            Inches(1.9),
            Inches(3.9),
            [0.55, 0.45],
            12,
        )
        estimates = [k.replace("_", " ") for k, v in scenario["assumptions"].items() if v.get("origin") == "estimate"]
        if estimates:
            deck.text(
                s,
                "Model estimates, not sourced: " + ", ".join(estimates),
                Inches(8.9),
                Inches(4.4),
                Inches(3.9),
                Inches(2),
                10,
                MUTED,
            )

    # 4. Scores
    factors = decision.get("factors") or data.get("factors", [])
    if factors:
        s = deck.slide(f"How the score is built: {decision.get('overall', rec.get('overall', ''))} / 100", "Scoring")
        rows = [["Factor", "Weight", "Score", "Evidence", "Claims"]]
        rows += [
            [
                f["label"],
                f"{f['weight'] * 100:.0f}%",
                str(f["score"]),
                f"{f['coverage'] * 100:.0f}%",
                ", ".join(f["claims"]) or "none",
            ]
            for f in factors
        ]
        deck.table(s, rows, Inches(0.85), Inches(1.9), Inches(11.6), [0.34, 0.12, 0.12, 0.14, 0.28], 13)
        deck.text(
            s,
            "Factors with no supporting evidence are held at 50. "
            "Weights may have been adjusted in the decision workspace.",
            Inches(0.85),
            Inches(5.9),
            Inches(11.6),
            Inches(0.6),
            11,
            MUTED,
        )

    # 5. Competitors
    if data.get("competitors"):
        s = deck.slide("Competitors named in the evidence", "Market")
        rows = [["Competitor", "Positioning", "Price", "Evidence"]]
        rows += [
            [c["name"], c["positioning"], c.get("price_note") or "—", ", ".join(c["claims"])]
            for c in data["competitors"][:8]
        ]
        deck.table(s, rows, Inches(0.85), Inches(1.9), Inches(11.6), [0.2, 0.5, 0.15, 0.15], 13)

    # 6. SWOT
    swot = data.get("swot")
    if swot:
        s = deck.slide("SWOT", "Strategic frameworks")
        for i, k in enumerate(("strengths", "weaknesses", "opportunities", "threats")):
            x = Inches(0.85 + (i % 2) * 6.1)
            y = Inches(1.8 + (i // 2) * 2.55)
            deck.text(s, k.upper(), x, y, Inches(5.8), Inches(0.35), 11, GOLD, bold=True)
            deck.text(
                s,
                "",
                x,
                y + Inches(0.35),
                Inches(5.8),
                Inches(2.1),
                13,
                INK,
                bullets=[f"{it['text']} ({', '.join(it['claims']) or 'no evidence'})" for it in swot.get(k, [])[:4]],
            )

    # 7. Risks and gaps
    s = deck.slide("Open questions and missing data", "Critique")
    deck.text(
        s,
        "",
        Inches(0.85),
        Inches(1.9),
        Inches(6),
        Inches(4.8),
        13,
        INK,
        bullets=[f"[{f['severity'].upper()}] {f['text']}" for f in data.get("findings", [])[:8]],
    )
    deck.text(
        s,
        "",
        Inches(7.2),
        Inches(1.9),
        Inches(5.5),
        Inches(4.8),
        13,
        INK,
        bullets=[f"{m['item']} ({m['impact']}): {m['plan']}" for m in data.get("missing", [])[:6]],
    )

    # 8. Evidence
    s = deck.slide("Evidence behind the analysis", "Evidence")
    deck.text(
        s,
        "",
        Inches(0.85),
        Inches(1.9),
        Inches(3.5),
        Inches(3),
        15,
        INK,
        bullets=[f"{STATUS_LABEL[k]}: {counts.get(k, 0)}" for k in STATUS_LABEL],
    )
    sources = {x["id"]: x for x in data.get("sources", [])}
    lines = []
    for c in [c for c in data.get("claims", []) if c["status"] == "supported"][:6]:
        src = next((sources.get(e["source_id"]) for e in c["evidence"] if e["relation"] == "supports"), None)
        lines.append(f"{c['key']} {c['text']}" + (f" ({src['publisher']})" if src else ""))
    deck.text(
        s,
        "",
        Inches(4.6),
        Inches(1.9),
        Inches(8.2),
        Inches(4.8),
        13,
        INK,
        bullets=lines or ["No claim is fully supported yet."],
    )

    # 9. Notice
    s = deck.slide("Important notice", "About this deck", dark=True)
    deck.text(s, DISCLAIMER, Inches(0.85), Inches(1.9), Inches(11.5), Inches(3), 16, IVORY)

    buf = io.BytesIO()
    deck.prs.save(buf)
    return buf.getvalue()
